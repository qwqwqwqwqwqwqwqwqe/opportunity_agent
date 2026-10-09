"""Synthetic documents only; fixtures never contain a real person's resume."""
import io
import json
import time
import zipfile

from opportunity_agent.resume_models import ParsedResumeDocument, ResumeBlock, ResumeDraft, ResumeFact, ResumeExperience

def document(text="GPA: 3.9/4.0\nTOEFL: 105\nSynthetic University Computer Science"):
    return ParsedResumeDocument(blocks=[ResumeBlock(block_id="p1", text=text, locator="段落 1")])

def draft():
    return ResumeDraft(facts=[
        ResumeFact(field="school", value="Synthetic University", confidence=.93, evidence="Synthetic University", block_ids=["p1"]),
        ResumeFact(field="gpa_raw", raw_value="GPA: 3.9/4.0", normalized_value=3.9, confidence=.98, evidence="GPA: 3.9/4.0", block_ids=["p1"]),
        ResumeFact(field="gpa_scale", value=4.0, confidence=.98, evidence="GPA: 3.9/4.0", block_ids=["p1"]),
    ], experiences=[ResumeExperience(kind="research", name="Signal Project", period="2026", role="复现实验",
        methods="Python, PyTorch", outcomes="精度 90%，延迟下降 10%", evidence="synthetic", block_ids=["p1"])])

class Extractor:
    mode, error = "llm", None
    def generate(self, parsed):
        return draft()

def wait_job(service, sid, job_id, status="review"):
    deadline = time.monotonic() + 12
    while time.monotonic() < deadline:
        job = service.get(sid, job_id)
        if job["status"] == status:
            return job
        if job["status"] == "failed" and status != "failed":
            raise AssertionError(job["error"])
        time.sleep(.02)
    raise AssertionError(f"job not {status}: {job}")

def zip_bytes(members):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return out.getvalue()

def docx_bytes():
    from docx import Document
    doc = Document()
    doc.add_paragraph("合成测试简历 Synthetic University | Computer Science")
    doc.add_paragraph("GPA: 3.9/4.0")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "科研信号分类"
    table.cell(1, 0).text = "方法与成果"
    table.cell(1, 1).text = "Python, PyTorch；精度 90%，延迟下降 10%"
    doc.add_paragraph("TOEFL: 105; 邮箱: synthetic@example.com; 手机: 13800138000")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()

def pdf_bytes(lines=None, pages=1, two_columns=False):
    # Tiny PDF fixture generator, including an explicit Unicode character map.
    # Its purpose is parser testing, not a user-facing document artifact.
    lines = lines if lines is not None else ["Synthetic University GPA: 3.9/4.0", "TOEFL: 105"]
    chars = sorted(set("".join(lines)))
    mappings = "\n".join(f"<{ord(c):04x}> <{ord(c):04x}>" for c in chars)
    cmap = ("/CIDInit /ProcSet findresource begin 12 dict begin begincmap "
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def "
            "/CMapName /Test def /CMapType 2 def 1 begincodespacerange <0000> <ffff> endcodespacerange "
            f"{len(chars)} beginbfchar\n{mappings}\nendbfchar endcmap CMapName currentdict /CMap defineresource pop end end").encode()
    def stream(data):
        return f"<< /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream"
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"",
        b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /Identity-H /DescendantFonts [4 0 R] /ToUnicode 5 0 R >>",
        b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light /CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 0 >> /DW 600 >>", stream(cmap)]
    kids = []
    for _ in range(pages):
        page_id = len(objs) + 1
        kids.append(f"{page_id} 0 R")
        content = []
        for index, line in enumerate(lines):
            x = 310 if two_columns and index % 2 else 40
            y = 750 - (index // 2 if two_columns else index) * 22
            content.append(f"BT /F1 10 Tf 1 0 0 1 {x} {y} Tm <{line.encode('utf-16-be').hex()}> Tj ET")
        objs += [f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {page_id+1} 0 R >>".encode(),
                 stream("\n".join(content).encode())]
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {pages} >>".encode()
    out = b"%PDF-1.4\n"; offsets = [0]
    for index, obj in enumerate(objs, 1):
        offsets.append(len(out)); out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    out += f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out

